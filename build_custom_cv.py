import os
import time
import base64
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
import pymupdf

# ==========================================
# 1. GENERATE DOCX FILES
# ==========================================
def create_docx(version_type, output_path):
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(9.5)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(2)
    normal_style.paragraph_format.space_before = Pt(0)
    
    def add_p(text='', bold=False, italic=False, font_size=9.5, color_rgb=(30, 41, 59), align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(*color_rgb)
        return p

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(15, 23, 42)
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="8" w:space="2" w:color="0F172A"/></w:pBdr>')
        pPr.append(pBdr)
        return p

    def add_item_header(left_title, right_meta, subtitle='', tech=''):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        
        r1 = p.add_run(left_title)
        r1.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(15, 23, 42)
        
        if right_meta:
            r_meta = p.add_run(f'  |  {right_meta}')
            r_meta.bold = False
            r_meta.font.size = Pt(9.5)
            r_meta.font.color.rgb = RGBColor(71, 85, 105)
            
        if subtitle or tech:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_before = Pt(0)
            p2.paragraph_format.space_after = Pt(2)
            p2.paragraph_format.keep_with_next = True
            if subtitle:
                r_sub = p2.add_run(subtitle)
                r_sub.bold = True
                r_sub.font.size = Pt(9.5)
                r_sub.font.color.rgb = RGBColor(2, 132, 199)
            if subtitle and tech:
                r_sep = p2.add_run(' — ')
                r_sep.font.size = Pt(9.5)
                r_sep.font.color.rgb = RGBColor(100, 116, 139)
            if tech:
                r_tech = p2.add_run(f'Tech: {tech}')
                r_tech.italic = True
                r_tech.font.size = Pt(9)
                r_tech.font.color.rgb = RGBColor(71, 85, 105)

    def add_bullet(lead_bold, rest_text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25)
        if lead_bold:
            r_bold = p.add_run(lead_bold)
            r_bold.bold = True
            r_bold.font.size = Pt(9.5)
            r_bold.font.color.rgb = RGBColor(15, 23, 42)
        r_rest = p.add_run(rest_text)
        r_rest.font.size = Pt(9.5)
        r_rest.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # Header
    add_p('MOHAMED MOSTAFA HELAL', bold=True, font_size=18, color_rgb=(15, 23, 42), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=1)
    headline_text = 'AI & Machine Learning Engineer | LLMs, Agentic Systems & Computer Vision'
    add_p(headline_text, bold=True, font_size=10.5, color_rgb=(2, 132, 199), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=3)
    
    if version_type == 'Egypt':
        contact_line_1 = 'Giza, Egypt  |  +20 103 046 9659  |  mohamed200031921@gmail.com  |  Military Status: Exempted'
    else:
        contact_line_1 = 'Egypt (Remote, Flexible Timezone)  |  +20 103 046 9659  |  mohamed200031921@gmail.com'
    contact_line_2 = 'LinkedIn: linkedin.com/in/mohamed-mostafa-454b70271  |  GitHub: github.com/DarkStrox  |  Hugging Face: huggingface.co/AitroniXAI'
    
    add_p(contact_line_1, font_size=9, color_rgb=(71, 85, 105), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=1)
    add_p(contact_line_2, font_size=9, color_rgb=(71, 85, 105), align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)

    # Summary
    add_section_header('Professional Summary')
    summary_text = (
        'AI & Machine Learning Engineer and 2026 Computer Science graduate (GPA 3.62/4.0, Honors) specialized in LLM agentic systems, edge inference, and real-time computer vision. '
        'Developed AitroniX 1s at an AI startup under senior engineering direction—an open-source 2.6B bilingual reasoning SLM achieving 68.75% strict AST match on BFCL v4, optimized for mobile edge deployment via ~1.56 GB GGUF. '
        'Proven hands-on track record building local hybrid RAG architectures (FAISS + BM25) and deploying hardware-accelerated multi-stream video analytics pipelines running at 30+ FPS.'
    )
    add_p(summary_text, font_size=9.5, space_before=2, space_after=4)

    # Experience
    add_section_header('Professional Experience')
    
    # Aitronix
    add_item_header('Aitronix (AI Startup)', 'Jan 2026 – Present | Giza, Egypt' if version_type == 'Egypt' else 'Jan 2026 – Present | Remote', subtitle='AI Engineer')
    add_bullet('AitroniX 1s SLM Development: ', 'Collaborated under Senior AI Engineer to fine-tune and release AitroniX 1s, an open-source 2.6B bilingual (Arabic/English) reasoning SLM on a hybrid Conv+Attention architecture with 128K context using Unsloth and LoRA/QLoRA.')
    add_bullet('Structured Tool-Calling Evaluation: ', 'Benchmarked structured tool calling on Berkeley Function-Calling Leaderboard (BFCL v4, 400 problems), achieving 68.75% strict AST match and 100% valid tool syntax (vs. Qwen3.5-9B at 60.13% and Gemma-4-8B at 48.39%).')
    add_bullet('Edge Quantization & Model Serving: ', 'Quantized model weights to Q4_K_M GGUF (~1.56 GB footprint) for private mobile edge execution; validated production API serving endpoints using vLLM and SGLang.')
    add_bullet('Real-Time Video Analytics: ', 'Engineered hardware-accelerated computer vision pipelines and multi-camera RTSP streaming engines, sustaining 30+ FPS across concurrent live feeds via OpenCL GPU acceleration and asynchronous worker queues.')

    # Drug Control Fund
    add_item_header('Drug Control & Addiction Prevention Fund', 'Jan 2024 – Dec 2025 | Cairo, Egypt' if version_type == 'Egypt' else 'Jan 2024 – Dec 2025 | Egypt', subtitle="Full-Stack Developer — Mutatawi' Rasmi")
    add_bullet('End-to-End Platform Development: ', "Developed the Mutatawi' Rasmi web application ecosystem from scratch, implementing responsive user interfaces, backend business logic, and relational database schemas to coordinate nationwide anti-drug youth initiatives.")
    add_bullet('Geospatial Hotspot Analytics: ', 'Engineered data-driven geospatial hotspot analytics and risk visualization maps across Egyptian governorates, empowering administrators to track field missions and allocate operational resources.')
    add_bullet('Volunteer Operations & Engagement: ', 'Implemented volunteer engagement engines, automated leaderboards, milestone verification workflows, and mission tracking modules to streamline volunteer coordination and reporting.')

    # Samsung Innovation Campus
    add_item_header('Samsung Innovation Campus', 'Jul 2024 – Oct 2024 | Giza, Egypt' if version_type == 'Egypt' else 'Jul 2024 – Oct 2024 | Egypt', subtitle='AI & Data Science Resident Intern')
    add_bullet('Quantitative Feature Engineering: ', 'Co-developed a quantitative time-series forecasting framework by processing raw market data and engineering 15 financial features, including rolling standard deviations, momentum indices, and historical lags.')
    add_bullet('Predictive Modeling & Optimization: ', 'Constructed and tuned ensemble regression and time-series models (XGBoost, Random Forest, and SARIMA), systematically reducing Mean Absolute Percentage Error (MAPE) from 8.3% to 6.5%.')
    add_bullet('Exploratory Data Analysis & Dashboards: ', 'Conducted extensive exploratory data analysis (EDA) on customer retention datasets using Pandas and designed interactive dashboards with Plotly and Seaborn to track operational attrition trends.')

    # Projects
    add_section_header('Technical Projects')

    # Project 1: AitroniX 1s
    add_item_header('AitroniX 1s — Bilingual 2.6B Agentic Reasoning SLM', 'Jan 2026', tech='Hybrid Conv+Attention, Unsloth, LoRA/QLoRA, GGUF, vLLM, SGLang, BFCL v4')
    add_bullet('Architecture & Open-Source Release: ', 'Engineered an open-source 2.6B bilingual (Arabic/English) reasoning model featuring a hybrid Convolution + Attention architecture, 128K context window, and native <think> Chain-of-Thought reasoning.')
    add_bullet('Benchmark Performance: ', 'Evaluated on Berkeley Function-Calling Leaderboard (BFCL v4 AST, all 400 problems), achieving 68.75% strict AST match with 100% valid tool syntax (vs. Qwen3.5-9B at 60.13%, Gemma-4-8B at 48.39%, and Gemma-4-2B at 35.98%).')
    add_bullet('Evaluation & Quantization: ', 'Benchmarked 77.83 on ToolSandbox, 85.49% on IFStruct, 51.87% on AIME25 math, and quantized weights to Q4_K_M GGUF (~1.56 GB) for zero-latency private inference on mobile edge devices.')

    # Project 2: HyperRAG
    add_item_header('HyperRAG — Local Hybrid RAG Engine', '2026', tech='FAISS, BM25, Reciprocal Rank Fusion (RRF), llama.cpp, Gemma 4 E4B, Python')
    add_bullet('Hybrid Retrieval Pipeline: ', 'Engineered a local hybrid retrieval pipeline combining dense vector search (FAISS) and sparse lexical keyword search (BM25) via Reciprocal Rank Fusion (RRF), eliminating keyword entity retrieval misses.')
    add_bullet('Private Offline Inference: ', 'Deployed local LLM generation via llama.cpp using Gemma 4 E4B for private, zero-cloud inference, featuring dynamic 800-token chunking (120-token overlap) and verifiable document citation tracking.')

    # Project 3: Face Recognition Engine
    add_item_header('Aitronix Real-Time High-FPS Face Recognition Engine', '2026', tech='Python, OpenCV, YuNet, SFace, OpenCL, Multithreading, RTSP, LibJPEG')
    add_bullet('High-Speed Face Recognition: ', 'Built a real-time face recognition pipeline utilizing YuNet for single-pass tensor face detection (~2.5ms inference via OpenCL GPU acceleration) and SFace for 128D embeddings with Cosine Similarity vector matching.')
    add_bullet('Multithreading & Byte Decoding: ', 'Designed non-blocking asynchronous multithreading with decoupled worker queues delivering sustained 30+ FPS; implemented custom native LibJPEG byte decoder to eliminate network buffering delay over RTSP and USB feeds.')

    # Project 4: CamScence
    add_item_header('CamScence — Multi-Camera Edge Traffic Analytics System', '2026', tech='YOLO (YOLO26n), RTSP Multi-Streaming, Speed Estimation, Spatial Heatmaps, Python')
    add_bullet('Multi-Camera Edge Pipeline: ', 'Architected an edge video analytics system concurrently ingesting 11+ RTSP IP camera streams with dynamic multi-grid switching and automated feed monitoring.')
    add_bullet('Tracking & Speed Estimation: ', 'Implemented real-time vehicle detection, persistent multi-object tracking, automated vehicle speed estimation (km/h), and spatial heatmap overlays for traffic congestion and bottleneck analysis.')

    # Project 5: Surveillance Threat Detection
    add_item_header('Real-Time Video Surveillance Threat & Burglary Detection', '2025 – 2026', tech='Microsoft Florence-2, OpenAI CLIP, OpenCV, Zero-Shot Grounding, Python')
    add_bullet('Multimodal Video Analysis: ', 'Developed an automated video surveillance pipeline combining Vision-Language Models (Microsoft Florence-2, OpenAI CLIP) and OpenCV to detect security incidents, burglary, and physical violence.')
    add_bullet('Zero-Shot Alerting: ', 'Applied zero-shot visual classification and phrase grounding to score threat probabilities frame-by-frame and trigger dynamic real-time on-screen security alerts.')

    # Project 6: YOLO-World
    add_item_header('Open-Vocabulary Real-Time Detection & Tracking (YOLO-World)', '2025 – 2026', tech='YOLO-World, Open-Vocabulary Vision, PyTorch, Real-Time GPU Tracking')
    add_bullet('Promptable Zero-Shot Tracking: ', 'Deployed YOLO-World to detect and track custom arbitrary natural language prompted visual concepts without model retraining or additional labeled data.')
    add_bullet('High-FPS GPU Execution: ', 'Optimized local GPU inference pipelines to execute real-time object tracking at 70+ FPS on high-speed video feeds while maintaining continuous spatial ID tracking.')

    # Project 7: Arabian Horses Graduation Project
    add_item_header('Integrated AI Ecosystem for Arabian Horses (Graduation Project)', 'Sep 2025 – Jun 2026', tech='React.js, ASP.NET Core Web API, SQL Server, Python, ConvNeXt, PyTorch')
    add_bullet('Full-Stack Platform: ', 'Architected an end-to-end multi-role web platform (React frontend, ASP.NET Core REST API backend, SQL Server) for horse auctions, farm management (Marabet), and medical record tracking.')
    add_bullet('Deep Learning Pedigree & Valuation: ', 'Integrated a Python deep learning pipeline utilizing ConvNeXt for horse pedigree and strain identification alongside regression models for market valuation and price estimation (Grade: Excellent).')

    # Other Projects
    p_other = doc.add_paragraph()
    p_other.paragraph_format.space_before = Pt(3)
    p_other.paragraph_format.space_after = Pt(3)
    r_other_lbl = p_other.add_run('Additional Projects: ')
    r_other_lbl.bold = True
    r_other_lbl.font.size = Pt(9)
    r_other_lbl.font.color.rgb = RGBColor(15, 23, 42)
    r_other_txt = p_other.add_run('NLP Benchmark (DistilBERT 91.36% on 50k IMDb reviews vs RNN/LSTM baselines with FP16); WhatsApp AI Assistant (n8n + Supabase pgvector + WhatsApp Cloud API); Moondream Edge VLM (webcam VQA); Plant Disease Classifier (38-class CNN, 95.4% accuracy).')
    r_other_txt.font.size = Pt(9)
    r_other_txt.font.color.rgb = RGBColor(71, 85, 105)

    # Skills
    add_section_header('Technical Skills')
    def add_skill_line(cat, skills):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r_cat = p.add_run(f'{cat}: ')
        r_cat.bold = True
        r_cat.font.size = Pt(9.5)
        r_cat.font.color.rgb = RGBColor(15, 23, 42)
        r_sk = p.add_run(skills)
        r_sk.font.size = Pt(9.5)
        r_sk.font.color.rgb = RGBColor(51, 65, 85)

    add_skill_line('LLMs & Agentic Systems', 'LLM Fine-Tuning (Unsloth, LoRA/QLoRA), Structured Tool Calling (JSON/XML, AST), Function Calling, GGUF Quantization, vLLM, SGLang, llama.cpp, Ollama, Hugging Face Transformers, Tokenization, Prompt Engineering')
    add_skill_line('Retrieval & Search (RAG)', 'Hybrid RAG, Dense Vector Search (FAISS), Sparse Lexical Indexing (BM25), Reciprocal Rank Fusion (RRF), Vector Databases (pgvector / Supabase)')
    add_skill_line('Computer Vision & Multimodal', 'YOLO (v8, YOLO-World, YOLO26n), Vision-Language Models (Florence-2, CLIP, Moondream), Face Recognition (YuNet, SFace), ConvNeXt, 3D CNNs, MobileNet, Transfer Learning, OpenCV, OpenCL GPU Acceleration, RTSP Multi-Streaming, Spatial Heatmap Tracking & Speed Estimation')
    add_skill_line('Machine Learning & Data Science', 'PyTorch, TensorFlow, Keras, Scikit-learn, XGBoost, Random Forest, SARIMA, Time-Series Forecasting, Feature Engineering, Pandas, NumPy, Plotly, Seaborn')
    add_skill_line('Backend & Engineering', 'Python, C#, ASP.NET Core Web API, React.js, JavaScript, PostgreSQL, Microsoft SQL Server, RESTful APIs, Asynchronous Multithreading, n8n Automation, Docker, Git, GitHub')
    add_skill_line('Languages', 'Arabic (Native), English (Professional / Fluent)')

    # Education
    add_section_header('Education')
    add_item_header('Misr University for Science and Technology (MUST)', 'Giza, Egypt | Sep 2022 – Jun 2026', subtitle='Bachelor of Science in Computer Science, Major in Artificial Intelligence')
    add_bullet('Academic Distinction: ', 'Cumulative GPA: 3.62 / 4.0 (Excellent with Honors)')
    add_bullet('Graduation Project: ', 'Integrated AI Ecosystem for Arabian Horses (Grade: Excellent)')

    # Certifications
    add_section_header('Certifications & Professional Credentials')
    add_bullet('Samsung Innovation Campus Certified Professional in AI & Data Science (2024): ', 'Issued by Samsung Electronics & Ministry of Communications and Information Technology (MCIT).')
    add_bullet('Deep Learning & Computer Vision Specialization: ', 'Comprehensive study in CNN architectures, object detection, and transfer learning.')
    add_bullet('Machine Learning & Advanced Statistical Modeling: ', 'Time-series forecasting, regression, and ensemble models.')

    doc.save(output_path)
    print(f'Saved DOCX: {output_path}')


# ==========================================
# 2. GENERATE HTML & COMPILE TO PDF VIA SELENIUM
# ==========================================
def generate_html_content(version_type):
    if version_type == 'Egypt':
        contact_line_1 = """
        <span class="contact-item">Giza, Egypt</span>
        <span class="contact-sep">•</span>
        <span class="contact-item"><a href="tel:+201030469659">+20 103 046 9659</a></span>
        <span class="contact-sep">•</span>
        <span class="contact-item"><a href="mailto:mohamed200031921@gmail.com">mohamed200031921@gmail.com</a></span>
        <span class="contact-sep">•</span>
        <span class="contact-item">Military Status: Exempted</span>
        """
        exp_loc_aitronix = "Giza, Egypt"
        exp_loc_fund = "Cairo, Egypt"
        exp_loc_samsung = "Giza, Egypt"
    else:
        contact_line_1 = """
        <span class="contact-item">Egypt (Remote, Flexible Timezone)</span>
        <span class="contact-sep">•</span>
        <span class="contact-item"><a href="tel:+201030469659">+20 103 046 9659</a></span>
        <span class="contact-sep">•</span>
        <span class="contact-item"><a href="mailto:mohamed200031921@gmail.com">mohamed200031921@gmail.com</a></span>
        """
        exp_loc_aitronix = "Remote"
        exp_loc_fund = "Egypt"
        exp_loc_samsung = "Egypt"

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>Mohamed Mostafa Helal - Curriculum Vitae</title>
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600&display=swap" rel="stylesheet">
  <style>
    @page {{
      size: A4;
      margin: 7mm 11mm 7mm 11mm;
    }}
    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }}
    body {{
      font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
      color: #1e293b;
      background-color: #ffffff;
      font-size: 8.35pt;
      line-height: 1.32;
    }}
    a {{
      color: #0284c7;
      text-decoration: none;
    }}
    .page {{
      width: 100%;
      position: relative;
    }}
    .page-break {{
      page-break-before: always;
      break-before: page;
    }}
    .header {{
      text-align: center;
      padding-bottom: 5px;
      margin-bottom: 6px;
      border-bottom: 2px solid #0f172a;
    }}
    .name {{
      font-size: 19pt;
      font-weight: 800;
      letter-spacing: 0.6px;
      color: #0f172a;
      text-transform: uppercase;
      margin-bottom: 2px;
      line-height: 1.1;
    }}
    .headline {{
      font-size: 9.3pt;
      font-weight: 700;
      color: #0284c7;
      margin-bottom: 4px;
      letter-spacing: 0.2px;
    }}
    .contact-info {{
      display: flex;
      justify-content: center;
      align-items: center;
      flex-wrap: wrap;
      gap: 7px;
      font-size: 8pt;
      color: #475569;
      font-weight: 500;
      margin-bottom: 2px;
    }}
    .contact-item {{
      display: inline-flex;
      align-items: center;
      gap: 3px;
    }}
    .contact-sep {{
      color: #cbd5e1;
      font-size: 7.5pt;
    }}
    .section {{
      margin-bottom: 6.5px;
    }}
    .section-title {{
      font-size: 9.3pt;
      font-weight: 800;
      text-transform: uppercase;
      letter-spacing: 0.5px;
      color: #0f172a;
      border-bottom: 1.2px solid #0f172a;
      padding-bottom: 2px;
      margin-bottom: 4px;
      display: flex;
      align-items: center;
      gap: 5px;
    }}
    .section-title .bar {{
      width: 3.5px;
      height: 10.5px;
      background-color: #0284c7;
      display: inline-block;
      border-radius: 1px;
    }}
    .entry {{
      margin-bottom: 4.5px;
    }}
    .entry:last-child {{
      margin-bottom: 0;
    }}
    .entry-header-row1 {{
      display: flex;
      justify-content: space-between;
      align-items: baseline;
    }}
    .entry-header-row2 {{
      display: flex;
      justify-content: space-between;
      align-items: baseline;
      margin-top: 0.5px;
      margin-bottom: 1.5px;
    }}
    .entry-title {{
      font-size: 8.7pt;
      font-weight: 700;
      color: #0f172a;
    }}
    .entry-subtitle {{
      font-size: 8.1pt;
      font-weight: 600;
      color: #0284c7;
    }}
    .entry-tech {{
      font-size: 7.6pt;
      font-weight: 600;
      color: #475569;
      font-style: italic;
    }}
    .entry-meta {{
      font-size: 7.8pt;
      font-weight: 600;
      color: #64748b;
      white-space: nowrap;
    }}
    .entry-bullets {{
      list-style-type: none;
      padding-left: 0;
      margin-top: 1px;
    }}
    .entry-bullets li {{
      position: relative;
      padding-left: 10px;
      margin-bottom: 1.5px;
      font-size: 8.1pt;
      color: #334155;
      line-height: 1.32;
      text-align: justify;
    }}
    .entry-bullets li:last-child {{
      margin-bottom: 0;
    }}
    .entry-bullets li::before {{
      content: "▪";
      position: absolute;
      left: 0;
      top: -0.5px;
      color: #0284c7;
      font-size: 8.2pt;
    }}
    .highlight {{
      font-weight: 700;
      color: #0f172a;
    }}
    .metric-badge {{
      display: inline-block;
      background: #e0f2fe;
      color: #0369a1;
      font-weight: 700;
      padding: 0 3px;
      border-radius: 3px;
      font-size: 7.6pt;
    }}
    code.inline-code {{
      font-family: 'JetBrains Mono', monospace;
      font-size: 7.3pt;
      color: #0369a1;
      background: #f1f5f9;
      padding: 0 2px;
      border-radius: 2px;
    }}
    .skills-container {{
      display: flex;
      flex-direction: column;
      gap: 2.5px;
    }}
    .skill-row {{
      display: flex;
      align-items: baseline;
      font-size: 8.1pt;
      line-height: 1.32;
    }}
    .skill-category {{
      font-weight: 700;
      color: #0f172a;
      min-width: 142px;
      flex-shrink: 0;
    }}
    .skill-list {{
      color: #334155;
    }}
    .cert-item {{
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
      font-size: 8.1pt;
      background: #f8fafc;
      padding: 4.5px 8px;
      border-radius: 3px;
      border-left: 3px solid #0284c7;
      border-top: 1px solid #e2e8f0;
      border-right: 1px solid #e2e8f0;
      border-bottom: 1px solid #e2e8f0;
      margin-bottom: 3.5px;
    }}
    .cert-item:last-child {{
      margin-bottom: 0;
    }}
  </style>
</head>
<body>

  <!-- ==================== PAGE 1 ==================== -->
  <div class="page">
    <header class="header">
      <div class="name">Mohamed Mostafa Helal</div>
      <div class="headline">AI & Machine Learning Engineer | LLMs, Agentic Systems & Computer Vision</div>
      <div class="contact-info">
        {contact_line_1}
      </div>
      <div class="contact-info">
        <span class="contact-item">LinkedIn: <a href="https://linkedin.com/in/mohamed-mostafa-454b70271" target="_blank">linkedin.com/in/mohamed-mostafa-454b70271</a></span>
        <span class="contact-sep">•</span>
        <span class="contact-item">GitHub: <a href="https://github.com/DarkStrox" target="_blank">github.com/DarkStrox</a></span>
        <span class="contact-sep">•</span>
        <span class="contact-item">Hugging Face: <a href="https://huggingface.co/AitroniXAI" target="_blank">huggingface.co/AitroniXAI</a></span>
      </div>
    </header>

    <!-- Professional Summary -->
    <section class="section">
      <div class="section-title"><span class="bar"></span>Professional Summary</div>
      <p style="font-size: 8.15pt; color: #334155; line-height: 1.34; text-align: justify;">
        AI & Machine Learning Engineer and 2026 Computer Science graduate (GPA 3.62/4.0, Honors) specialized in LLM agentic systems, edge inference, and real-time computer vision. Developed AitroniX 1s at an AI startup under senior engineering direction—an open-source 2.6B bilingual reasoning SLM achieving <span class="highlight">68.75% strict AST match on BFCL v4</span>, optimized for mobile edge deployment via ~1.56 GB GGUF. Proven hands-on track record building local hybrid RAG architectures (FAISS + BM25) and deploying hardware-accelerated multi-stream video analytics pipelines running at <span class="highlight">30+ FPS</span>.
      </p>
    </section>

    <!-- Professional Experience -->
    <section class="section">
      <div class="section-title"><span class="bar"></span>Professional Experience</div>
      
      <!-- Aitronix -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Aitronix (AI Startup)</span>
          <span class="entry-meta">{exp_loc_aitronix}</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-subtitle">AI Engineer</span>
          <span class="entry-meta">Jan 2026 – Present</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">AitroniX 1s Model Fine-Tuning:</span> Collaborated under Senior AI Engineer to fine-tune and release AitroniX 1s, an open-source 2.6B bilingual (Arabic/English) reasoning SLM on a hybrid Conv+Attention architecture with 128K context using Unsloth and LoRA/QLoRA.</li>
          <li><span class="highlight">Structured Tool-Calling Evaluation:</span> Benchmarked structured tool calling on Berkeley Function-Calling Leaderboard (BFCL v4, 400 problems), achieving a <span class="highlight">68.75% strict AST match</span> and 100% valid tool syntax (vs. Qwen3.5-9B at 60.13% and Gemma-4-8B at 48.39%).</li>
          <li><span class="highlight">Edge Quantization & Model Serving:</span> Quantized model weights into Q4_K_M GGUF (~1.56 GB memory footprint) for private on-device mobile execution and configured high-throughput server deployment using vLLM and SGLang.</li>
          <li><span class="highlight">Real-Time Video Analytics:</span> Engineered hardware-accelerated computer vision pipelines and multi-camera RTSP streaming engines, sustaining <span class="highlight">30+ FPS</span> across concurrent live feeds via OpenCL GPU acceleration and asynchronous worker queues.</li>
        </ul>
      </div>

      <!-- Drug Control Fund -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Drug Control & Addiction Prevention Fund</span>
          <span class="entry-meta">{exp_loc_fund}</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-subtitle">Full-Stack Developer — Mutatawi' Rasmi (متطوع رسمي)</span>
          <span class="entry-meta">Jan 2024 – Dec 2025</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">End-to-End System Development:</span> Developed the Mutatawi' Rasmi web ecosystem from scratch, implementing responsive user interfaces, backend business logic, and relational database schemas to coordinate nationwide anti-drug youth initiatives.</li>
          <li><span class="highlight">Geospatial Hotspot Analytics:</span> Engineered data-driven geospatial hotspot analytics and risk visualization maps across Egyptian governorates, empowering administrators to track field missions and allocate operational resources.</li>
          <li><span class="highlight">Volunteer Operations & Engagement:</span> Implemented volunteer engagement modules, automated leaderboards, milestone verification workflows, and mission tracking modules to manage volunteer participation.</li>
        </ul>
      </div>

      <!-- Samsung Innovation Campus -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Samsung Innovation Campus</span>
          <span class="entry-meta">{exp_loc_samsung}</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-subtitle">AI & Data Science Resident Intern</span>
          <span class="entry-meta">Jul 2024 – Oct 2024</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Quantitative Feature Engineering:</span> Co-developed a quantitative time-series forecasting framework by processing raw market data and engineering 15 financial features, including rolling standard deviations, momentum indices, and historical lags.</li>
          <li><span class="highlight">Predictive Modeling & Optimization:</span> Constructed and tuned ensemble regression and time-series models (XGBoost, Random Forest, and SARIMA), systematically reducing Mean Absolute Percentage Error (MAPE) from <span class="highlight">8.3% to 6.5%</span>.</li>
          <li><span class="highlight">Exploratory Data Analysis & Dashboards:</span> Conducted extensive exploratory data analysis (EDA) on customer retention datasets using Pandas and designed interactive dashboards with Plotly and Seaborn to track operational attrition trends.</li>
        </ul>
      </div>
    </section>

    <!-- Technical Projects (Part 1) -->
    <section class="section">
      <div class="section-title"><span class="bar"></span>Featured Technical Projects</div>

      <!-- AitroniX 1s -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">AitroniX 1s: Bilingual 2.6B Agentic Reasoning SLM</span>
          <span class="entry-meta"><a href="https://huggingface.co/AitroniXAI/AitroniXAi-1s-GGUF" target="_blank">Hugging Face Model</a> | Jan 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: Hybrid Conv+Attention Architecture, GGUF (Q4_K_M), Unsloth, BFCL v4 AST, vLLM, SGLang, llama.cpp</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Architecture & Open-Source Release:</span> Engineered an open-source 2.6B bilingual (Arabic/English) reasoning model featuring a hybrid Conv+Attention architecture, 128K context window, and native <code class="inline-code">&lt;think&gt;</code> Chain-of-Thought reasoning.</li>
          <li><span class="highlight">Benchmark Evaluation:</span> Evaluated on Berkeley Function-Calling Leaderboard (BFCL v4 AST, 400 problems), achieving <span class="highlight">68.75% strict AST match</span> with 100% valid tool syntax (vs. Qwen3.5-9B at 60.13%, Gemma-4-8B at 48.39%, and Gemma-4-2B at 35.98%).</li>
          <li><span class="highlight">Edge Inference:</span> Benchmarked 77.83 on ToolSandbox, 85.49% on IFStruct, 51.87% on AIME25 math, and quantized weights to Q4_K_M GGUF (~1.56 GB) for zero-latency private inference on mobile edge devices.</li>
        </ul>
      </div>

      <!-- HyperRAG -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">HyperRAG: Local Hybrid RAG Engine (BM25 + FAISS with Gemma 4 E4B)</span>
          <span class="entry-meta">Generative AI & Search | 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: FAISS, BM25, Reciprocal Rank Fusion (RRF), llama.cpp, Gemma 4 E4B, Python</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Hybrid Retrieval Pipeline:</span> Engineered a local hybrid retrieval pipeline combining dense vector search (<span class="highlight">FAISS</span>) with sparse lexical keyword search (<span class="highlight">BM25</span>) via <span class="highlight">Reciprocal Rank Fusion (RRF)</span>, eliminating keyword entity retrieval misses.</li>
          <li><span class="highlight">Private Offline Inference:</span> Deployed local LLM generation via llama.cpp using Gemma 4 E4B for private, zero-cloud inference, featuring dynamic 800-token chunking (120-token overlap) and verifiable document citation tracking.</li>
        </ul>
      </div>

      <!-- Face Recognition -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Aitronix Real-Time High-FPS Face Recognition Engine</span>
          <span class="entry-meta">Computer Vision & Systems | 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: Python, OpenCV, YuNet, SFace, OpenCL GPU Acceleration, Multithreading, RTSP, LibJPEG</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">High-Speed Face Recognition:</span> Built a real-time face recognition pipeline utilizing YuNet for single-pass tensor face detection (~2.5ms inference via OpenCL GPU acceleration) and SFace for 128D embeddings with Cosine Similarity vector matching.</li>
          <li><span class="highlight">Multithreading & Byte Decoding:</span> Designed non-blocking asynchronous multithreading with decoupled worker queues delivering sustained <span class="highlight">30+ FPS</span>; implemented custom native LibJPEG byte decoder to eliminate network buffering delay over RTSP and USB feeds.</li>
        </ul>
      </div>
    </section>
  </div>

  <!-- ==================== PAGE 2 ==================== -->
  <div class="page page-break">
    <!-- Technical Projects (Part 2) -->
    <section class="section">
      <div class="section-title"><span class="bar"></span>Featured Technical Projects (Continued)</div>

      <!-- CamScence -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">CamScence: Multi-Camera Edge Traffic Analytics System</span>
          <span class="entry-meta">Video Analytics & Edge Vision | 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: YOLO (YOLO26n), RTSP Multi-Streaming (11+ Cameras), Speed Estimation, Spatial Heatmaps, Python</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Multi-Camera Edge Infrastructure:</span> Architected an edge video analytics system concurrently ingesting <span class="highlight">11+ RTSP IP camera streams</span> with dynamic multi-grid switching and automated feed monitoring.</li>
          <li><span class="highlight">Tracking & Speed Estimation:</span> Implemented real-time vehicle detection, persistent multi-object tracking, automated vehicle speed estimation (<span class="highlight">km/h</span>), and spatial heatmap overlays for traffic congestion and bottleneck analysis.</li>
        </ul>
      </div>

      <!-- Surveillance Threat Detection -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Real-Time Video Surveillance Threat & Burglary Detection</span>
          <span class="entry-meta">Vision-Language Models & Security | 2025 – 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: Microsoft Florence-2, OpenAI CLIP, OpenCV, Zero-Shot Grounding, Python</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Multimodal Video Analysis:</span> Developed an automated video surveillance pipeline combining Vision-Language Models (<span class="highlight">Microsoft Florence-2, OpenAI CLIP</span>) and OpenCV to detect security incidents, burglary, and physical violence.</li>
          <li><span class="highlight">Zero-Shot Alerting:</span> Applied zero-shot visual classification and phrase grounding to score threat probabilities frame-by-frame and trigger dynamic real-time on-screen security alerts.</li>
        </ul>
      </div>

      <!-- YOLO-World -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Open-Vocabulary Real-Time Detection & Tracking (YOLO-World)</span>
          <span class="entry-meta">Computer Vision & Zero-Shot AI | 2025 – 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: YOLO-World, Open-Vocabulary Vision, PyTorch, Real-Time GPU Tracking</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Promptable Zero-Shot Tracking:</span> Deployed YOLO-World to detect and track custom arbitrary natural language prompted visual concepts without model retraining or additional labeled data.</li>
          <li><span class="highlight">High-FPS GPU Execution:</span> Optimized local GPU inference pipelines to execute real-time object tracking at <span class="highlight">70+ FPS</span> on high-speed video feeds while maintaining continuous spatial ID tracking.</li>
        </ul>
      </div>

      <!-- Arabian Horses -->
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Integrated AI Ecosystem for Arabian Horses (Graduation Project)</span>
          <span class="entry-meta">Full-Stack & AI Engineer | Sep 2025 – Jun 2026</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-tech">Tech: React.js, ASP.NET Core REST API, SQL Server, Python, ConvNeXt, PyTorch</span>
        </div>
        <ul class="entry-bullets">
          <li><span class="highlight">Full-Stack Platform:</span> Architected an end-to-end multi-role web platform (React frontend, ASP.NET Core REST API backend, SQL Server) for horse auctions, farm management (Marabet), and medical record tracking.</li>
          <li><span class="highlight">Deep Learning Pedigree & Valuation:</span> Integrated a Python deep learning pipeline utilizing <span class="highlight">ConvNeXt</span> for horse pedigree and strain identification alongside regression models for market valuation and price estimation (<span class="metric-badge">Grade: Excellent</span>).</li>
        </ul>
      </div>

      <!-- Compressed Other Projects -->
      <div style="font-size: 7.9pt; color: #475569; margin-top: 3.5px; padding: 3px 6px; background: #f8fafc; border-radius: 3px; border: 1px solid #e2e8f0;">
        <strong style="color: #0f172a;">Additional Projects:</strong> NLP Benchmark (DistilBERT 91.36% on 50k IMDb reviews vs RNN/LSTM baselines with FP16); Dynamic WhatsApp AI Recruitment Assistant (n8n + Supabase pgvector + WhatsApp Cloud API); Moondream Edge VLM (webcam VQA); Plant Disease Classifier (38-class CNN, 95.4% accuracy).
      </div>
    </section>

    <!-- Technical Skills -->
    <section class="section">
      <div class="section-title"><span class="bar"></span>Technical Skills</div>
      <div class="skills-container">
        <div class="skill-row">
          <span class="skill-category">LLMs & Agentic Systems:</span>
          <span class="skill-list"><span class="highlight">LLM Fine-Tuning (Unsloth, LoRA/QLoRA)</span>, Structured Tool Calling (JSON/XML, AST), Function Calling, GGUF Quantization, vLLM, SGLang, llama.cpp, Ollama, Hugging Face Transformers, Tokenization, Prompt Engineering</span>
        </div>
        <div class="skill-row">
          <span class="skill-category">Retrieval & Search (RAG):</span>
          <span class="skill-list"><span class="highlight">Hybrid RAG</span>, Dense Vector Search (FAISS), Sparse Lexical Indexing (BM25), Reciprocal Rank Fusion (RRF), Vector Databases (pgvector / Supabase)</span>
        </div>
        <div class="skill-row">
          <span class="skill-category">Computer Vision & Multimodal:</span>
          <span class="skill-list"><span class="highlight">YOLO (v8, YOLO-World, YOLO26n)</span>, Vision-Language Models (Florence-2, CLIP, Moondream), Face Recognition (YuNet, SFace), ConvNeXt, 3D CNNs, MobileNet, Transfer Learning, OpenCV, OpenCL GPU Acceleration, RTSP Multi-Streaming, Spatial Heatmaps</span>
        </div>
        <div class="skill-row">
          <span class="skill-category">Machine Learning & Analytics:</span>
          <span class="skill-list"><span class="highlight">PyTorch, TensorFlow, Keras</span>, Scikit-learn, XGBoost, Random Forest, SARIMA, Time-Series Forecasting, Feature Engineering, Pandas, NumPy, Plotly, Seaborn</span>
        </div>
        <div class="skill-row">
          <span class="skill-category">Backend & Engineering:</span>
          <span class="skill-list"><span class="highlight">Python, C#, ASP.NET Core Web API</span>, React.js, JavaScript, PostgreSQL, Microsoft SQL Server, RESTful APIs, Asynchronous Multithreading, n8n Automation, Docker, Git, GitHub</span>
        </div>
        <div class="skill-row">
          <span class="skill-category">Languages:</span>
          <span class="skill-list">Arabic (Native / Fluent), English (Professional / Fluent)</span>
        </div>
      </div>
    </section>

    <!-- Education -->
    <section class="section">
      <div class="section-title"><span class="bar"></span>Education</div>
      <div class="entry">
        <div class="entry-header-row1">
          <span class="entry-title">Misr University for Science and Technology (MUST)</span>
          <span class="entry-meta">Giza, Egypt</span>
        </div>
        <div class="entry-header-row2">
          <span class="entry-subtitle">Bachelor of Science in Computer Science, Major in Artificial Intelligence</span>
          <span class="entry-meta">Graduated: June 2026</span>
        </div>
        <div style="font-size: 8.1pt; color: #334155; display: flex; gap: 15px;">
          <span><span class="highlight">Cumulative GPA:</span> <span class="metric-badge">3.62 / 4.0</span> (Excellent with Honors)</span>
          <span><span class="highlight">Graduation Project:</span> Integrated AI Ecosystem for Arabian Horses (<span class="metric-badge">Grade: Excellent</span>)</span>
        </div>
      </div>
    </section>

    <!-- Certifications -->
    <section class="section" style="margin-bottom: 0;">
      <div class="section-title"><span class="bar"></span>Certifications & Professional Credentials</div>
      <div class="cert-item">
        <div>
          <strong style="color: #0f172a; font-size: 8.2pt;">Samsung Innovation Campus Certified Professional in Artificial Intelligence & Data Science</strong>
          <div style="color: #64748b; font-size: 7.7pt; margin-top: 0.5px;">Issued by Samsung Electronics & Ministry of Communications and Information Technology (MCIT) • Machine Learning, Deep Learning, Statistical Modeling, and Computer Vision</div>
        </div>
        <div style="font-weight: 700; color: #0284c7; font-size: 8pt; white-space: nowrap; margin-left: 10px;">Issued: 2024</div>
      </div>
      <div class="cert-item">
        <div>
          <strong style="color: #0f172a; font-size: 8.2pt;">Deep Learning & Computer Vision Specialization</strong>
          <div style="color: #64748b; font-size: 7.7pt; margin-top: 0.5px;">Comprehensive study in Convolutional Neural Networks (CNNs), Object Detection, and Transfer Learning Architectures</div>
        </div>
        <div style="font-weight: 700; color: #0284c7; font-size: 8pt; white-space: nowrap; margin-left: 10px;">Credential</div>
      </div>
      <div class="cert-item">
        <div>
          <strong style="color: #0f172a; font-size: 8.2pt;">Machine Learning & Advanced Statistical Modeling</strong>
          <div style="color: #64748b; font-size: 7.7pt; margin-top: 0.5px;">Time-series forecasting, quantitative feature engineering, regression, and tree-based ensemble models</div>
        </div>
        <div style="font-weight: 700; color: #0284c7; font-size: 8pt; white-space: nowrap; margin-left: 10px;">Credential</div>
      </div>
    </section>
  </div>

</body>
</html>
"""
    return html

def render_html_to_pdf(html_content, output_html_path, output_pdf_path):
    with open(output_html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
        
    chrome_options = Options()
    chrome_options.add_argument("--headless=new")
    chrome_options.add_argument("--disable-gpu")
    chrome_options.add_argument("--no-sandbox")
    chrome_options.add_argument("--disable-dev-shm-usage")
    chrome_options.add_argument("--allow-file-access-from-files")
    
    driver = webdriver.Chrome(options=chrome_options)
    file_url = "file:///" + os.path.abspath(output_html_path).replace("\\", "/")
    print(f"Loading {file_url} in Chrome...")
    driver.get(file_url)
    time.sleep(2.0)
    
    print_options = {
        'landscape': False,
        'displayHeaderFooter': False,
        'printBackground': True,
        'preferCSSPageSize': True,
        'marginTop': 0,
        'marginBottom': 0,
        'marginLeft': 0,
        'marginRight': 0,
    }
    
    result = driver.execute_cdp_cmd("Page.printToPDF", print_options)
    pdf_bytes = base64.b64decode(result['data'])
    with open(output_pdf_path, "wb") as f:
        f.write(pdf_bytes)
    driver.quit()
    print(f"Successfully rendered PDF: {output_pdf_path}")
    
    # Check page count
    doc = pymupdf.open(output_pdf_path)
    print(f"Total Pages in {os.path.basename(output_pdf_path)}: {len(doc)}")
    return len(doc)

def render_pdf_to_images(pdf_path, output_dir, prefix):
    doc = pymupdf.open(pdf_path)
    img_paths = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=200)
        img_path = os.path.join(output_dir, f"{prefix}_page_{i+1}.png")
        pix.save(img_path)
        img_paths.append(img_path)
        print(f"Saved preview: {img_path}")
    return img_paths

if __name__ == '__main__':
    preview_dir = "e:/Aitronix Projects/cv_preview"
    os.makedirs(preview_dir, exist_ok=True)
    
    # 1. Generate DOCX
    create_docx('Egypt', os.path.join(preview_dir, 'Mohamed_Mostafa_Helal_CV_Egypt.docx'))
    create_docx('International', os.path.join(preview_dir, 'Mohamed_Mostafa_Helal_CV_International.docx'))
    
    # 2. Generate Egypt PDF
    egypt_html = generate_html_content('Egypt')
    egypt_html_path = "e:/Aitronix Projects/cv_egypt.html"
    egypt_pdf_path = os.path.join(preview_dir, "Mohamed_Mostafa_Helal_CV_Egypt.pdf")
    pages_eg = render_html_to_pdf(egypt_html, egypt_html_path, egypt_pdf_path)
    render_pdf_to_images(egypt_pdf_path, preview_dir, "cv_egypt")
    
    # 3. Generate International PDF
    intl_html = generate_html_content('International')
    intl_html_path = "e:/Aitronix Projects/cv_international.html"
    intl_pdf_path = os.path.join(preview_dir, "Mohamed_Mostafa_Helal_CV_International.pdf")
    pages_intl = render_html_to_pdf(intl_html, intl_html_path, intl_pdf_path)
    render_pdf_to_images(intl_pdf_path, preview_dir, "cv_intl")

    # Update 2-page default
    import shutil
    shutil.copyfile(egypt_pdf_path, "e:/Aitronix Projects/Mohamed_Mostafa_Helal_CV_2Page.pdf")
    shutil.copyfile(egypt_pdf_path, os.path.join(preview_dir, "Mohamed_Mostafa_Helal_CV v7.pdf"))
    print("Done! Both PDFs and previews generated.")
